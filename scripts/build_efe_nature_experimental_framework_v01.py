from __future__ import annotations

import argparse
import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


PARAGRAPH_REPLACEMENTS: dict[int, tuple[str, str]] = {
    5: (
        "研究主图设计 · 3–4 年路线",
        "实验论文框架 · PRL 发布后应用路线",
    ),
    6: (
        "心内膜弹性纤维增生症（EFE）",
        "心内膜弹性纤维增生症（EFE）",
    ),
    7: (
        "面向 Nature 的六张主图与逐 Panel 实验设计",
        "面向 Nature 的六张主图与逐 Panel 实验设计",
    ),
    8: (
        "中心命题：发育期异常血流形成可累积的机械记忆，通过心内膜-ECM 正反馈，将可逆修复性纤维化锁定为不可逆的 EFE。",
        "中心命题（待实验验证）：发育期异常血流/负荷通过可累积的局部力学历史驱动谱系特异的心内膜-ECM 反馈，使可逆修复跨入持续性 fibroelastic remodeling；冻结的 PRL 框架用于前瞻定位，而不替代因果实验。",
    ),
    9: ("优势组合", "双论文分工"),
    10: (
        "斑马鱼活体成像 × 三维/四维重建 × 理论力学 × 生信与跨物种分析",
        "PRL 冻结力学框架 × EFE 4D 成像 × 谱系/多组学 × 因果干预 × 人源验证",
    ),
    11: ("内部研究设计稿｜2026年8月", "内部实验论文框架｜v01｜2026年8月"),
    13: ("论文叙事骨架", "双论文关系与证据防火墙"),
    14: (
        "现象发现（4D成像）→ 个体化重建（力学历史）→ 盲预测（相图与阈值）→ 细胞来源（谱系）→ 分子门控（多组学与因果）→ 人类疾病验证。",
        "PRL 预先发表/冻结 → EFE 表型建立 → 个体化力学映射 → EFE-Blind 前瞻预测 → 谱系与细胞状态 → Gate X 因果干预 → 人源验证。",
    ),
    15: ("术语和边界", "论文身份与术语边界"),
    16: (
        "在斑马鱼完成跨物种验证前，使用“EFE-like fibroelastic remodeling”，不直接宣称建立了人 EFE。",
        "PRL 是理论先行的独立论文；EFE Nature 是实验疾病机制论文，不作为 PRL 的应用附录，也不承担 PRL 的普适理论证明。",
    ),
    17: (
        "疾病表型必须同时包含心内膜下定位、纤维弹性 ECM、持续性和功能后果；仅有 col1a/postn 上升不构成 EFE。",
        "EFE 只调用已发表/冻结的 PRL FrameworkVersion；允许校准几何适配器和实验量级，不允许用 EFE 终点回写核心方程、无量纲组或预测门。",
    ),
    18: (
        "成年斑马鱼损伤/再生模型主要作为“可逆修复性纤维化”参照，主疾病模型应放在发育期心脏。",
        "跨物种验证完成前使用“EFE-like fibroelastic remodeling”；疾病表型必须同时具备心内膜下定位、纤维弹性 ECM、持续性和功能后果。",
    ),
    19: (
        "细胞来源允许为时间依赖的双来源或多来源，不预设 EndMT 与心外膜来源必须二选一。",
        "成年斑马鱼损伤/再生只作为可逆修复参照；细胞来源允许时间依赖的双来源或多来源，不预设 EndMT 与心外膜来源必须二选一。",
    ),
    20: ("全项目统计和证据原则", "四类数据分区与证据原则"),
    21: (
        "鱼是生物学重复单位；图像块、切片、细胞和材料点属于嵌套观测，不得充当独立 n。",
        "鱼是生物学重复单位；图像块、切片、细胞和材料点属于嵌套观测，不得充当独立 n。",
    ),
    22: (
        "模型校准组与盲测组在看到病理终点前划分；主要预测、误差阈值和失败标准预先登记。",
        "`PRL-Cal` 只估计几何/材料量级；`PRL-Val` 检验冻结的通用力学预测；`EFE-Discovery` 用于疾病发现；`EFE-Blind` 从开始即封存。",
    ),
    23: (
        "每个关键结论至少由两种正交技术支持，例如活体信号 + 终点组织学、谱系 + 空间原位、结构救援 + 功能救援。",
        "同一鱼、时相或处理组不得同时承担校准与盲验证；FrameworkVersion、GeometryAdapter、预测方向、误差阈值和失败标准必须在 EFE-Blind 解盲前登记。",
    ),
    24: (
        "主图只保留支持中心命题的结果；完整 QC、阴性结果、替代模型和通路筛选进入 Extended Data。",
        "每个关键结论至少由两种正交技术支持；主图保留中心证据，完整 QC、阴性预测、替代模型、版本变化和通路筛选进入 Extended Data。",
    ),
    66: (
        "Figure 2｜从 4D 图像重建个体化力学历史并提出相变模型",
        "Figure 2｜调用冻结 PRL 框架重建个体化力学历史",
    ),
    67: (
        "主张：EFE-like 状态由可累积的局部机械暴露和 ECM 反馈决定，模型能够输出每个心内膜位置的状态转换风险，而非只给出全心平均参数。",
        "主张：预先发表/冻结的 PRL 框架可把 EFE 个体的 4D 运动与流动转换为逐位置机械历史和前瞻风险量；该映射提出可证伪预测，但不单独证明疾病因果。",
    ),
    69: (
        "Figure 2 成图原型｜由 4D 成像重建个体化机械历史，并提出 Normal–Reversible–Locked 状态转变模型。",
        "Figure 2 成图原型｜以冻结的 PRL FrameworkVersion 为核心，由 4D 成像重建个体化机械历史并生成可检验风险图。",
    ),
    75: (
        "B｜运动边界与流场重建的实验校准",
        "B｜GeometryAdapter 与边界条件的实验校准",
    ),
    76: (
        "展示内容：成像获得的壁运动、血细胞轨迹与模拟速度矢量叠加。",
        "展示内容：成像获得的壁运动、血细胞轨迹与计算速度/牵引矢量叠加，并标注 PRL 核心层与 EFE 适配层。",
    ),
    77: (
        "实验/分析：使用一部分个体校准边界条件、黏度和入口/出口约束；用未用于求解的流速、时相和轨迹做验证。",
        "实验/分析：使用 `PRL-Cal`/EFE 发现数据估计几何映射、黏度和入口/出口量级；核心方程与无量纲结构不重拟合；用 `PRL-Val` 检验未参与校准的速度、相位和轨迹。",
    ),
    78: (
        "决定性读出：模型能重建局部速度和时相特征，而非仅匹配心率或平均输出。",
        "决定性读出：适配层在独立个体中重建局部速度和时相特征；若失败，先报告 ModelBoundary，不用疾病终点补偿性调参。",
    ),
    87: ("E｜最小心内膜-ECM-flow 状态模型", "E｜冻结框架与 EFE 应用层接口"),
    88: (
        "展示内容：模型框图包含心内膜状态、ECM 沉积/降解、组织刚度和流动反馈；显示可能的正反馈与恢复路径。",
        "展示内容：框图列出 FrameworkVersion、GeometryAdapter、MechanicalHistory、ProspectiveRule 和 ModelBoundary 五个接口，并区分固定项与可校准项。",
    ),
    89: (
        "实验/分析：先使用可辨识的最小变量集；把分子门控写成待实验确定的 Gate X，避免在模型建立初期堆叠通路。",
        "实验/分析：PRL 核心方程、无量纲组和适用域保持冻结；EFE 只加入个体几何、观测误差与实验可测量量的适配。分子门控继续记为待实验确定的 Gate X。",
    ),
    90: (
        "决定性读出：模型在参数范围内产生可逆修复态与锁定 fibroelastosis 态两个不同吸引子。",
        "决定性读出：每个心内膜位置获得带不确定性的机械历史和候选风险量；任何“阈值/吸引子/迟滞”主张留给 Figure 3 的实验比较与替代模型检验。",
    ),
    91: ("F｜训练数据与保留验证数据的预先划分", "F｜四类数据分区与冻结顺序"),
    92: (
        "展示内容：明确哪些扰动、剂量和时间点用于参数校准，哪些从一开始封存为盲测。",
        "展示内容：以流程图明确 `PRL-Cal → PRL-Val → EFE-Discovery → EFE-Blind` 的用途、锁定时间和访问权限。",
    ),
    93: (
        "实验/分析：按鱼而非图像块/细胞进行训练-验证划分；记录先验、参数后验和不可辨识参数。",
        "实验/分析：按鱼而非图像块/细胞划分；同一扰动批次不跨校准与盲测；记录数据清单、访问时间、版本、先验和不可辨识参数。",
    ),
    94: (
        "决定性读出：模型选择和参数估计在看到盲测终点前冻结，避免后验调参。",
        "决定性读出：FrameworkVersion、适配器、主要终点和失败判据在 EFE-Blind 病理终点解盲前冻结。",
    ),
    95: ("G｜机械剂量-ECM 刚度相图", "G｜机械历史到候选实验风险图"),
    96: (
        "展示内容：二维或三维相图显示正常、可逆纤维化和锁定 EFE-like 区域，以及跨越阈值的路径。",
        "展示内容：候选状态图显示正常、可逆重塑和持续 EFE-like 区域，并给出空间热点、暴露时间窗与不确定区间。",
    ),
    97: (
        "实验/分析：进行参数扫描、稳定性分析和不确定性传播；标出实际鱼样本在相图中的位置。",
        "实验/分析：先比较连续剂量、阈值、积分暴露和迟滞候选模型，再做不确定性传播；样本位置只用于生成前瞻预测，不提前定性为状态跃迁。",
    ),
    98: (
        "决定性读出：相图给出非线性阈值、发育窗口和恢复路径差异三个可实验检验的结论。",
        "决定性读出：形成空间热点、发育窗口和恢复路径差异三个可实验检验的候选结论，并明确各自的失败模式。",
    ),
    99: ("H｜预注册的三项盲预测", "H｜EFE-Blind 的三项前瞻预测"),
    100: (
        "展示内容：在主图末端列出下一张图将检验的预测：空间风险区、发育时间窗、相同平均流量但不同脉动的分叉结局。",
        "展示内容：在主图末端列出下一张图将检验的预测：空间风险区、发育时间窗、相同平均流量但不同脉动/相位的不同结局。",
    ),
    101: (
        "实验/分析：在封存数据前记录预测方向、允许误差和失败判据；不把敏感性分析当作预测成功。",
        "实验/分析：在 `EFE-Blind` 解盲前记录预测方向、允许误差、判分程序和失败判据；不把敏感性分析或事后版本当作预测成功。",
    ),
    102: (
        "决定性读出：形成明确的可证伪承诺，使理论在论文中承担发现作用。",
        "决定性读出：形成明确的可证伪承诺，使已发表的 PRL 框架在 EFE 论文中承担前瞻解释作用，同时保留失败结果。",
    ),
    103: (
        "Go/No-Go：模型必须在校准数据上通过独立流场验证，并给出至少三项方向明确、可在一年内检验的盲预测。",
        "Go/No-Go：FrameworkVersion 与适配器必须先通过 `PRL-Val`/独立流场验证，并在 EFE-Blind 解盲前给出至少三项方向明确、可在一年内判分的预测。",
    ),
    105: (
        "Figure 3｜盲法验证机械阈值、发育窗口与迟滞",
        "Figure 3｜前瞻盲法检验空间热点、发育窗口与持续性",
    ),
    106: (
        "主张：EFE-like 形成取决于力学波形、暴露时序和正反馈阈值；恢复平均血流并不总能逆转已锁定的状态。",
        "主张：冻结框架可在未见病理终点的条件下前瞻分层 EFE-like 的空间、时间窗和恢复结局；只有在替代模型被排除后，才讨论阈值或迟滞机制。",
    ),
    108: (
        "Figure 3 成图原型｜以前瞻性盲法实验验证空间热点、发育时窗、非线性阈值与滞后。",
        "Figure 3 成图原型｜以 EFE-Blind 前瞻实验检验空间热点、发育时窗、剂量非线性与持续性/迟滞候选。",
    ),
    110: ("A｜保留扰动矩阵与实验随机化", "A｜EFE-Blind 扰动矩阵与实验随机化"),
    118: ("C｜模型生成个体化空间风险图", "C｜冻结模型生成个体化空间风险图"),
    120: (
        "实验/分析：冻结模型后输入该鱼的几何与力学历史；预先规定风险区面积和峰值位置指标。",
        "实验/分析：冻结模型后输入该鱼的几何与力学历史；分析者不可访问组织学终点；预先规定风险区面积、峰值位置和不确定区间。",
    ),
    134: ("G｜恢复实验揭示迟滞", "G｜恢复实验检验持续性与迟滞候选"),
    137: (
        "决定性读出：阈值前可恢复、阈值后持续，支持 ECM-力学正反馈形成机械记忆。",
        "决定性读出：若阈值前可恢复、阈值后持续，且一般发育延迟、累积损伤和处理毒性不能解释，才支持 ECM-力学反馈与迟滞候选。",
    ),
    141: (
        "决定性读出：至少两项关键盲预测成功，且成功不能由平均流量或一般发育延迟模型解释。",
        "决定性读出：至少两类关键盲预测达到预登记标准，且成功不能由平均流量或一般发育延迟模型解释；阴性预测同样进入成绩单。",
    ),
    142: (
        "Go/No-Go：空间位置、发育窗口和恢复迟滞三类预测中至少两类在独立批次成立；否则停止扩展组学，先修订机制模型。",
        "Go/No-Go：空间位置、发育窗口和恢复结局三类预测中至少两类在独立批次达到预登记标准；否则保留负结果，停止扩大组学，并仅在下一版本前瞻修订模型。",
    ),
    269: ("Extended Data 分配建议", "Extended Data 分配建议"),
    270: (
        "原则：主图回答机制问题，Extended Data 证明结论可靠、边界清楚、替代解释被排除。",
        "原则：主图回答实验机制问题，Extended Data 证明 FrameworkVersion 可追溯、结果可靠、边界清楚且替代解释被排除。",
    ),
    274: (
        "Extended Data 4：流体与组织力学模型的收敛、敏感性、参数可辨识性和替代模型比较。",
        "Extended Data 4：PRL FrameworkVersion、GeometryAdapter、流体/组织力学收敛、参数可辨识性、盲预测版本和替代模型比较。",
    ),
    279: ("3–4 年并行排程建议", "PRL 冻结依赖下的 3–4 年并行排程"),
    280: ("最早需要落实的外部能力", "最早需要落实的外部能力与前置条件"),
    284: (
        "必要时引入局部组织力学测量合作，但其方法学必须服务于 Figure 2 的可测模型参数。",
        "局部组织力学测量必须服务于 Figure 2 的可测输入/验证量，并提前确定哪些数据属于 `PRL-Cal`、`PRL-Val` 或 EFE 专用分区。",
    ),
    286: ("工作参考文献", "工作参考文献（沿用原稿，投稿前逐条复核）"),
}


TABLE_REPLACEMENTS: dict[tuple[int, int, int], str] = {
    (0, 1, 1): "是否真正建立了满足四项疾病定义的斑马鱼 EFE-like 表型？",
    (0, 2, 1): "冻结的 PRL 框架能否从个体 4D 数据生成可追溯的局部机械历史？",
    (0, 3, 1): "EFE-Blind 是否支持空间、时间窗、剂量和恢复结局的前瞻预测？",
    (0, 4, 1): "哪些谱系在何时进入持续细胞状态，局部机械历史能否解释异质性？",
    (0, 5, 1): "Gate X 是否具有必要性、充分性和上位性，并可被时间特异救援？",
    (0, 6, 1): "同一状态和机制是否存在于人 EFE，并可在人源体系中因果重现？",
    (1, 1, 1): "建立 EFE-like 主模型；完成纵向 4D 成像、ECM 组成、持续性和功能判定。",
    (1, 1, 2): "登记 PRL FrameworkVersion 与 GeometryAdapter；完成 `PRL-Cal`/`PRL-Val` 分区和可测量量映射。",
    (1, 1, 3): "Figure 1 四项疾病标准成立；适配器通过独立验证；EFE-Blind 保持封存。",
    (1, 2, 1): "完成 EFE-Blind 正交扰动、恢复实验和双谱系先导。",
    (1, 2, 2): "冻结空间风险、发育窗口和恢复预测；执行盲法判分与替代模型比较。",
    (1, 2, 3): "Figure 3 至少两类前瞻预测达到预登记标准。",
    (1, 3, 1): "完成谱系、多组学、Gate X 因果与阈值前/后救援。",
    (1, 3, 2): "整合机械历史—细胞状态网络；只在前瞻版本中更新 Gate X 应用参数。",
    (1, 3, 3): "结构、细胞状态和心功能三层同步救援。",
    (1, 4, 1): "完成人组织、人源体系验证及独立病例/细胞系重复。",
    (1, 4, 2): "跨物种映射；冻结 EFE 机制版本、负结果、适用边界和可复算材料。",
    (1, 4, 3): "Figure 6 支撑保守机制；若人源因果不足则降低跨物种主张。",
}


INTERFACE_PAGE = [
    (
        "Heading 1",
        "PRL→EFE 冻结接口与可追溯要求",
    ),
    (
        "Figure Claim",
        "原则：EFE Nature 把 PRL 当作预先存在、可证伪的分析框架，而不是可随疾病终点反复调节的拟合器；疾病因果仍由实验独立建立。",
    ),
    ("List Bullet", "`FrameworkVersion`：记录论文/预印本版本、方程、无量纲组、适用域、失败边界、代码提交和 Figure 版本。"),
    ("List Bullet", "`GeometryAdapter`：规定 4D 分割、材料坐标、运动边界、流动/力学输入和测量误差如何进入冻结框架。"),
    ("List Bullet", "`MechanicalHistory`：输出逐鱼、逐位置的幅值、相位、牵引、耗散和记忆统计量，并保留不确定性。"),
    ("List Bullet", "`ProspectiveRule`：在 EFE-Blind 解盲前冻结空间热点、发育窗口、剂量与恢复结局的方向、阈值和判分程序。"),
    ("List Bullet", "`ModelBoundary`：记录框架失败、不可辨识、替代本构更优或实验暴露不可重建的条件；失败不被隐藏为参数更新。"),
    ("List Bullet", "因果证据：谱系来源、Gate X 必要性/充分性、早晚救援和人源重现均由 EFE 实验建立，不能由 PRL 风险图替代。"),
    (
        "Go NoGo",
        "Go/No-Go：只有当五项接口、四类数据分区和 EFE-Blind 判分程序均有版本号与时间戳时，Figure 2–3 才进入正式盲验证；否则仅作探索性分析。",
    ),
]


def replace_paragraph_text(paragraph, replacement: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = replacement
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.add_run(replacement)


def shade_paragraph(paragraph, fill: str, border: str | None = None) -> None:
    paragraph_properties = paragraph._p.get_or_add_pPr()
    shading = paragraph_properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        paragraph_properties.append(shading)
    shading.set(qn("w:fill"), fill)
    if border:
        paragraph_borders = paragraph_properties.find(qn("w:pBdr"))
        if paragraph_borders is None:
            paragraph_borders = OxmlElement("w:pBdr")
            paragraph_properties.append(paragraph_borders)
        left_border = OxmlElement("w:left")
        left_border.set(qn("w:val"), "single")
        left_border.set(qn("w:sz"), "18")
        left_border.set(qn("w:space"), "7")
        left_border.set(qn("w:color"), border)
        paragraph_borders.append(left_border)


def replace_cell_text(cell, replacement: str) -> None:
    first_paragraph = cell.paragraphs[0]
    replace_paragraph_text(first_paragraph, replacement)
    for paragraph in cell.paragraphs[1:]:
        replace_paragraph_text(paragraph, "")


def build_document(source: Path, output: Path) -> None:
    if not source.is_file():
        raise FileNotFoundError(source)
    if output.exists():
        raise FileExistsError(f"Refusing to overwrite existing output: {output}")
    if not output.parent.exists():
        output.parent.mkdir(parents=True, exist_ok=False)
    shutil.copy2(source, output)

    document = Document(output)
    original_paragraphs = list(document.paragraphs)
    for index, (expected, replacement) in PARAGRAPH_REPLACEMENTS.items():
        paragraph = original_paragraphs[index]
        if paragraph.text != expected:
            raise ValueError(
                f"Template paragraph {index} mismatch: expected {expected!r}, found {paragraph.text!r}"
            )
        replace_paragraph_text(paragraph, replacement)

    for (table_index, row_index, column_index), replacement in TABLE_REPLACEMENTS.items():
        replace_cell_text(document.tables[table_index].cell(row_index, column_index), replacement)

    # Keep the six-figure overview heading and its compact table on one page.
    original_paragraphs[25].paragraph_format.page_break_before = True

    extended_data_anchor = original_paragraphs[269]
    for style_name, text in INTERFACE_PAGE:
        inserted = extended_data_anchor.insert_paragraph_before(text, style=style_name)
        if style_name == "Heading 1" and text.startswith("PRL→EFE"):
            inserted.paragraph_format.page_break_before = True
        if style_name == "Go NoGo":
            shade_paragraph(inserted, "FFF4D6", "A07416")

    original_paragraphs[269].paragraph_format.page_break_before = True
    original_paragraphs[279].paragraph_format.page_break_before = True
    original_paragraphs[286].paragraph_format.page_break_before = True

    core_properties = document.core_properties
    core_properties.title = "EFE Nature 实验论文框架：PRL 冻结框架的下游应用"
    core_properties.subject = "斑马鱼 EFE-like 表型、4D 力学、谱系、Gate X 与人源验证的六图实验路线"
    core_properties.author = "项目组"
    core_properties.keywords = "EFE; zebrafish; 4D imaging; mechanobiology; lineage; blind validation; PRL framework"
    core_properties.comments = "Derived from the retained EFE Nature figure-plan reference; v01, 2026-08-27."

    document.save(output)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    arguments = parser.parse_args()
    build_document(arguments.source, arguments.output)


if __name__ == "__main__":
    main()
