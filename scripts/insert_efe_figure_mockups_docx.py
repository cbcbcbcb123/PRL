from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from PIL import Image


SOURCE_DOCX = Path(r"C:\Users\chenb\Desktop\EFE_Nature_主图小图设计方案.docx")
OUTPUT_BASENAME = "EFE_Nature_主图成图版"
OUTPUT_DIR = Path(r"C:\Users\chenb\Desktop")
FIGURE_DIR = Path(r"E:\Temp-Projects\PRL\figures\efe_nature_mockups")
DOCX_FIGURE_DIR = FIGURE_DIR / "docx_optimized"


SOURCE_FIGURES = {
    1: FIGURE_DIR / "Figure1_EFE_phenotype_concept.png",
    2: FIGURE_DIR / "Figure2_mechanical_history_model_concept.png",
    3: FIGURE_DIR / "Figure3_blind_validation_concept.png",
    4: FIGURE_DIR / "Figure4_lineage_cell_state_concept.png",
    5: FIGURE_DIR / "Figure5_causal_rescue_concept.png",
    6: FIGURE_DIR / "Figure6_human_translation_concept.png",
}


FIGURES = {
    figure_number: DOCX_FIGURE_DIR / f"{source_path.stem}_docx.jpg"
    for figure_number, source_path in SOURCE_FIGURES.items()
}


CAPTIONS = {
    1: "Figure 1 成图原型｜建立斑马鱼发育性 EFE-like 表型：从亚心内膜 fibroelastic 层到舒张与血流功能障碍。",
    2: "Figure 2 成图原型｜由 4D 成像重建个体化机械历史，并提出 Normal–Reversible–Locked 状态转变模型。",
    3: "Figure 3 成图原型｜以前瞻性盲法实验验证空间热点、发育时窗、非线性阈值与滞后。",
    4: "Figure 4 成图原型｜将内膜/外膜谱系、单细胞状态与局部机械历史统一到时空状态图谱。",
    5: "Figure 5 成图原型｜以 Gate X 的必要性、充分性、上位性和时间依赖性救援建立因果链。",
    6: "Figure 6 成图原型｜在人 EFE 组织和人源内膜流动平台中验证保守机制及其干预窗口。",
}


ALT_TEXTS = {
    1: "Conceptual multi-panel Figure 1 showing a zebrafish developmental EFE-like phenotype and associated structure, ECM, stiffness, function, flow, and persistence readouts.",
    2: "Conceptual multi-panel Figure 2 showing the 4D image-to-mesh pipeline, individualized mechanical histories, feedback model, train-holdout design, and phase diagram.",
    3: "Conceptual multi-panel Figure 3 showing blinded validation of spatial risk, developmental windows, nonlinear threshold, and hysteresis.",
    4: "Conceptual multi-panel Figure 4 showing lineage tracing, single-cell states, spatial mechanics mapping, Gate X network, and a dual-origin EFE model.",
    5: "Conceptual multi-panel Figure 5 showing Gate X mechanism, necessity, sufficiency, epistasis, early prevention, late rescue, and a shifted phase boundary.",
    6: "Conceptual multi-panel Figure 6 showing human EFE tissue, spatial and single-nucleus profiling, cross-species mapping, a human endocardial-on-chip system, and rescue.",
}


def next_output_path() -> Path:
    candidate = OUTPUT_DIR / f"{OUTPUT_BASENAME}.docx"
    if not candidate.exists():
        return candidate
    version = 2
    while True:
        candidate = OUTPUT_DIR / f"{OUTPUT_BASENAME}_v{version}.docx"
        if not candidate.exists():
            return candidate
        version += 1


def prepare_docx_images() -> None:
    DOCX_FIGURE_DIR.mkdir(parents=True, exist_ok=True)
    for figure_number, source_path in SOURCE_FIGURES.items():
        output_path = FIGURES[figure_number]
        if output_path.exists():
            continue
        with Image.open(source_path) as source_image:
            clean_rgb = source_image.convert("RGB")
            clean_rgb.save(
                output_path,
                format="JPEG",
                quality=95,
                subsampling=0,
                optimize=True,
                progressive=True,
                dpi=(240, 240),
            )


def add_paragraph_after(anchor, style: str | None = None):
    paragraph = anchor._parent.add_paragraph()
    anchor._p.addnext(paragraph._p)
    if style:
        paragraph.style = style
    return paragraph


def ensure_styles(document: Document) -> None:
    styles = document.styles
    if "Concept Figure Caption" not in styles:
        style = styles.add_style("Concept Figure Caption", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(8.5)
        style.font.italic = True
        style.font.color.rgb = RGBColor(69, 77, 87)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.space_before = Pt(3)
        style.paragraph_format.space_after = Pt(3)
        style.paragraph_format.keep_with_next = False

    if "Concept Figure Note" not in styles:
        style = styles.add_style("Concept Figure Note", WD_STYLE_TYPE.PARAGRAPH)
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(8)
        style.font.color.rgb = RGBColor(122, 53, 73)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = False


def set_picture_alt_text(inline_shape, title: str, description: str) -> None:
    doc_properties = inline_shape._inline.docPr
    doc_properties.set("title", title)
    doc_properties.set("descr", description)


def find_claim_and_first_panel(document: Document, figure_number: int):
    heading_prefix = f"Figure {figure_number}"
    paragraphs = document.paragraphs
    for index, paragraph in enumerate(paragraphs):
        if paragraph.text.startswith(heading_prefix):
            if index + 2 >= len(paragraphs):
                raise RuntimeError(f"Figure {figure_number} heading has no following claim and panel paragraphs")
            return paragraphs[index + 1], paragraphs[index + 2]
    raise RuntimeError(f"Could not locate heading for Figure {figure_number}")


def insert_figure_block(document: Document, figure_number: int, image_width) -> None:
    claim, first_panel = find_claim_and_first_panel(document, figure_number)
    first_panel.paragraph_format.page_break_before = True

    image_paragraph = add_paragraph_after(claim)
    image_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    image_paragraph.paragraph_format.space_before = Pt(7)
    image_paragraph.paragraph_format.space_after = Pt(2)
    image_paragraph.paragraph_format.keep_with_next = False
    picture = image_paragraph.add_run().add_picture(str(FIGURES[figure_number]), width=image_width)
    set_picture_alt_text(
        picture,
        f"Figure {figure_number} conceptual completed-paper mockup",
        ALT_TEXTS[figure_number],
    )

    caption = add_paragraph_after(image_paragraph, "Concept Figure Caption")
    caption.add_run(CAPTIONS[figure_number])

    note = add_paragraph_after(caption, "Concept Figure Note")
    note.paragraph_format.keep_with_next = False
    note.add_run("说明：该图用于规划完成论文后的信息组织与视觉结构；图内曲线、热图、UMAP、统计条和显微图均为概念性数据占位，不代表真实实验结果。")


def main() -> None:
    if not SOURCE_DOCX.exists():
        raise FileNotFoundError(SOURCE_DOCX)
    missing = [str(path) for path in SOURCE_FIGURES.values() if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing source figure files: " + "; ".join(missing))

    prepare_docx_images()
    missing = [str(path) for path in FIGURES.values() if not path.exists()]
    if missing:
        raise FileNotFoundError("Missing optimized figure files: " + "; ".join(missing))

    document = Document(str(SOURCE_DOCX))
    ensure_styles(document)

    section = document.sections[0]
    usable_width = section.page_width - section.left_margin - section.right_margin
    image_width = int(usable_width * 0.98)

    for figure_number in sorted(FIGURES):
        insert_figure_block(document, figure_number, image_width)

    document.core_properties.title = "EFE Nature 主图成图版：六张主图概念原型与小图设计"
    document.core_properties.subject = "Endocardial fibroelastosis paper figure mockups and panel-by-panel experimental plan"
    document.core_properties.comments = (
        "Contains ImageGen-created conceptual figure mockups. All visualized data are placeholders and must be replaced with verified experimental results."
    )

    output_path = next_output_path()
    document.save(str(output_path))
    print(output_path)


if __name__ == "__main__":
    main()
