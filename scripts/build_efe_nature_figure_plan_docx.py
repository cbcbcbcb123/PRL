from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(r"C:\Users\chenb\Desktop\EFE_Nature_主图小图设计方案.docx")

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(11, 37, 69)
MUTED = RGBColor(95, 105, 118)
GOLD = RGBColor(160, 116, 22)
WHITE = RGBColor(255, 255, 255)
CALL_OUT = "EEF4FA"
GATE_FILL = "FFF4D6"
LIGHT_BLUE = "E8EEF5"


def set_run_font(run, ascii_name="Calibri", east_asia="Microsoft YaHei", size=None,
                 color=None, bold=None, italic=None):
    run.font.name = ascii_name
    if run._element.rPr is None:
        run._element.get_or_add_rPr()
    fonts = run._element.rPr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        run._element.rPr.insert(0, fonts)
    fonts.set(qn("w:ascii"), ascii_name)
    fonts.set(qn("w:hAnsi"), ascii_name)
    fonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def shade_paragraph(paragraph, fill, border=None):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    if border:
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is None:
            p_bdr = OxmlElement("w:pBdr")
            p_pr.append(p_bdr)
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "7")
        left.set(qn("w:color"), border)
        p_bdr.append(left)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_fixed_geometry(table, widths_dxa, indent_dxa=120):
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths_dxa[idx]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_run_font(run, size=9, color=MUTED)


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    h1 = styles["Heading 1"]
    h1.font.name = "Calibri"
    h1._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = BLUE
    h1.paragraph_format.space_before = Pt(18)
    h1.paragraph_format.space_after = Pt(10)
    h1.paragraph_format.keep_with_next = True

    h2 = styles["Heading 2"]
    h2.font.name = "Calibri"
    h2._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = BLUE
    h2.paragraph_format.space_before = Pt(14)
    h2.paragraph_format.space_after = Pt(7)
    h2.paragraph_format.keep_with_next = True

    h3 = styles["Heading 3"]
    h3.font.name = "Calibri"
    h3._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    h3.font.size = Pt(12)
    h3.font.bold = True
    h3.font.color.rgb = DARK_BLUE
    h3.paragraph_format.space_before = Pt(10)
    h3.paragraph_format.space_after = Pt(5)
    h3.paragraph_format.keep_with_next = True

    for list_style_name in ("List Bullet", "List Number"):
        list_style = styles[list_style_name]
        list_style.font.name = "Calibri"
        list_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        list_style.font.size = Pt(10.8)
        list_style.paragraph_format.left_indent = Inches(0.375)
        list_style.paragraph_format.first_line_indent = Inches(-0.188)
        list_style.paragraph_format.space_after = Pt(4)
        list_style.paragraph_format.line_spacing = 1.25

    panel = styles.add_style("Panel Heading", WD_STYLE_TYPE.PARAGRAPH)
    panel.font.name = "Calibri"
    panel._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    panel.font.size = Pt(11.5)
    panel.font.bold = True
    panel.font.color.rgb = DARK_BLUE
    panel.paragraph_format.space_before = Pt(9)
    panel.paragraph_format.space_after = Pt(3)
    panel.paragraph_format.keep_with_next = True

    panel_body = styles.add_style("Panel Body", WD_STYLE_TYPE.PARAGRAPH)
    panel_body.font.name = "Calibri"
    panel_body._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    panel_body.font.size = Pt(10.4)
    panel_body.paragraph_format.space_before = Pt(0)
    panel_body.paragraph_format.space_after = Pt(2.5)
    panel_body.paragraph_format.line_spacing = 1.17

    caption = styles.add_style("Figure Claim", WD_STYLE_TYPE.PARAGRAPH)
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.font.size = Pt(11.3)
    caption.font.bold = True
    caption.font.color.rgb = INK
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(10)
    caption.paragraph_format.line_spacing = 1.2

    gate = styles.add_style("Go NoGo", WD_STYLE_TYPE.PARAGRAPH)
    gate.font.name = "Calibri"
    gate._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    gate.font.size = Pt(10.8)
    gate.font.bold = True
    gate.font.color.rgb = INK
    gate.paragraph_format.space_before = Pt(10)
    gate.paragraph_format.space_after = Pt(8)
    gate.paragraph_format.left_indent = Inches(0.12)
    gate.paragraph_format.right_indent = Inches(0.12)
    gate.paragraph_format.line_spacing = 1.18


def configure_sections(doc):
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        header = section.header
        p = header.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run("EFE · Nature 主图设计方案")
        set_run_font(r, size=8.5, color=MUTED, bold=True)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.paragraph_format.space_before = Pt(0)
        add_page_number(fp)


def add_labeled_paragraph(doc, label, text, style="Panel Body"):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.keep_together = True
    r1 = p.add_run(label)
    set_run_font(r1, size=10.4, color=INK, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.4)
    return p


def add_panel(doc, letter, title, content, methods, readout):
    p = doc.add_paragraph(style="Panel Heading")
    p.add_run(f"{letter}｜{title}")
    p.paragraph_format.keep_with_next = True
    content_p = add_labeled_paragraph(doc, "展示内容：", content)
    content_p.paragraph_format.keep_with_next = True
    methods_p = add_labeled_paragraph(doc, "实验/分析：", methods)
    methods_p.paragraph_format.keep_with_next = True
    readout_p = add_labeled_paragraph(doc, "决定性读出：", readout)
    readout_p.paragraph_format.keep_with_next = False


def add_callout(doc, text, fill=CALL_OUT, border="2E74B5"):
    p = doc.add_paragraph(style="Figure Claim")
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(10)
    p.add_run(text)
    shade_paragraph(p, fill, border)
    return p


def add_gate(doc, text):
    p = doc.add_paragraph(style="Go NoGo")
    p.add_run("Go/No-Go：" + text)
    shade_paragraph(p, GATE_FILL, "A07416")
    return p


FIGURES = [
    {
        "title": "Figure 1｜建立具有疾病定义的斑马鱼 EFE-like 表型",
        "claim": "主张：发育期特定血流/负荷异常可诱导定位于心内膜下、富含纤维弹性 ECM、解除刺激后仍持续并损害舒张功能的 EFE-like 重塑。",
        "panels": [
            ("A", "临床问题到斑马鱼模型的尺度对应", "左侧呈现人 EFE 的组织学定义和关键功能后果；右侧示意斑马鱼发育期心室、心内膜、心肌与流出道，明确哪些结构可比、哪些不可直接等同。", "绘制疾病-模型概念图，预先规定斑马鱼结果在完成跨物种验证前使用“EFE-like”表述。", "读者在第一幅图即理解：研究模拟的是心内膜下 fibroelastic remodeling，而不是一般性心肌损伤瘢痕。"),
            ("B", "发育时间轴与扰动窗口", "标出正常心室形成、血流建立、瓣膜/流出道成熟和 ECM 重塑阶段，并显示早、中、晚三个干预窗口。", "先导实验采用可分级、可恢复的三类候选扰动：剪切/黏度改变、部分收缩力改变、流出阻力或局部负荷改变；完全停跳仅作边界对照。", "确定一个存活率、整体发育和心率可接受，同时能形成持续心内膜下 ECM 的主模型。"),
            ("C", "活体 4D 成像捕获病变起点", "正常组与 EFE-like 组从干预前到病变形成的连续体积时间序列，突出最早出现异常的心内膜区域。", "高速光片/共聚焦成像；同步记录心内膜、心肌、血细胞/流动和候选 ECM 报告信号；同一条鱼重复随访。", "病变不是终点切片中的偶发现象，而是从可定位的细胞/组织事件逐步形成。"),
            ("D", "三维结构重建与心内膜下层定量", "原始体数据、分割结果、表面网格和层厚热图并排展示。", "盲法分割心腔、心内膜、心肌和沉积层；报告分割一致性、重建误差、层厚、体积和空间异质性。", "EFE-like 组出现局灶或弥漫性心内膜下增厚，且其位置可在个体间通过解剖坐标配准。"),
            ("E", "纤维弹性 ECM 的组成证据", "胶原、弹性纤维及基底膜/纤维连接蛋白相关成分的空间共定位，配合组织学和蛋白层验证。", "候选 readout 包括胶原成像/染色、弹性纤维染色或蛋白验证、ECM 转录原位检测；优先使用三维定量而非代表性切片。", "同时满足“心内膜下定位 + 胶原/弹性 ECM + 结构增厚”，排除仅由炎症或一般 collagen scar 解释。"),
            ("F", "组织力学状态的空间测量", "病变区与非病变区的局部刚度、顺应性或黏弹性地图，与 ECM 沉积层叠加。", "根据平台条件选择微压痕、Brillouin/光学弹性测量或反演式组织力学；在固定和活体读出之间进行交叉验证。", "EFE-like 区域的力学异常与 ECM 结构一致，并能作为理论模型的独立输入或验证量。"),
            ("G", "心功能与流动后果", "正常与 EFE-like 个体的腔室运动、舒张充盈、搏出量、流速和回流/涡旋模式。", "高速成像进行壁运动追踪和心功能定量；血细胞追踪或微粒子测速重建局部流动。", "表型不仅是分子标志升高，还造成可重复的舒张受限、腔室生长异常或流动重构。"),
            ("H", "撤除刺激后的持续性与疾病判定", "早期恢复、晚期恢复和持续扰动三组的纵向曲线；对比可逆适应与不可逆锁定。", "在预设时间撤除负荷或恢复血流，继续随访 ECM、心内膜厚度和功能；设置恢复时间匹配的未扰动对照。", "晚于某一时点恢复血流仍不能清除 fibroelastic layer，构成后续“机械记忆/迟滞”假说的现象基础。"),
        ],
        "gate": "只有当主模型同时满足心内膜下定位、纤维弹性 ECM、持续性和功能后果四项标准时，才进入大规模谱系和组学实验。",
    },
    {
        "title": "Figure 2｜从 4D 图像重建个体化力学历史并提出相变模型",
        "claim": "主张：EFE-like 状态由可累积的局部机械暴露和 ECM 反馈决定，模型能够输出每个心内膜位置的状态转换风险，而非只给出全心平均参数。",
        "panels": [
            ("A", "图像到计算域的可复现流水线", "展示原始 4D 图像经分割、时相配准、表面重建、网格化到计算域的全过程。", "每一步保留版本、参数和质量指标；建立自动化/半自动化流程并锁定盲法 QC 阈值。", "不同批次和不同操作者得到相近几何、壁运动与局部力学指标。"),
            ("B", "运动边界与流场重建的实验校准", "成像获得的壁运动、血细胞轨迹与模拟速度矢量叠加。", "使用一部分个体校准边界条件、黏度和入口/出口约束；用未用于求解的流速、时相和轨迹做验证。", "模型能重建局部速度和时相特征，而非仅匹配心率或平均输出。"),
            ("C", "心内膜表面的机械暴露地图", "沿心动周期显示壁面剪切、振荡剪切、环向/纵向应变、压力和曲率的时空分布。", "将所有物理量投影到统一的心内膜材料坐标；保留均值、峰值、相位差和周期累积量。", "最早 EFE-like 区域具有可识别的机械暴露组合，而非由单一“低流量”指标解释。"),
            ("D", "单细胞/材料点的机械历史账本", "示意每个被追踪心内膜细胞或表面材料点对应一条随发育时间变化的力学轨迹。", "通过 4D 配准和材料点迁移，将剪切、应变、曲率和邻近 ECM 状态累积为个体级机械历史。", "后续组学和细胞命运可直接关联其局部机械历史，避免用全心平均值解释细胞异质性。"),
            ("E", "最小心内膜-ECM-flow 状态模型", "模型框图包含心内膜状态、ECM 沉积/降解、组织刚度和流动反馈；显示可能的正反馈与恢复路径。", "先使用可辨识的最小变量集；把分子门控写成待实验确定的 Gate X，避免在模型建立初期堆叠通路。", "模型在参数范围内产生可逆修复态与锁定 fibroelastosis 态两个不同吸引子。"),
            ("F", "训练数据与保留验证数据的预先划分", "明确哪些扰动、剂量和时间点用于参数校准，哪些从一开始封存为盲测。", "按鱼而非图像块/细胞进行训练-验证划分；记录先验、参数后验和不可辨识参数。", "模型选择和参数估计在看到盲测终点前冻结，避免后验调参。"),
            ("G", "机械剂量-ECM 刚度相图", "二维或三维相图显示正常、可逆纤维化和锁定 EFE-like 区域，以及跨越阈值的路径。", "进行参数扫描、稳定性分析和不确定性传播；标出实际鱼样本在相图中的位置。", "相图给出非线性阈值、发育窗口和恢复路径差异三个可实验检验的结论。"),
            ("H", "预注册的三项盲预测", "在主图末端列出下一张图将检验的预测：空间风险区、发育时间窗、相同平均流量但不同脉动的分叉结局。", "在封存数据前记录预测方向、允许误差和失败判据；不把敏感性分析当作预测成功。", "形成明确的可证伪承诺，使理论在论文中承担发现作用。"),
        ],
        "gate": "模型必须在校准数据上通过独立流场验证，并给出至少三项方向明确、可在一年内检验的盲预测。",
    },
    {
        "title": "Figure 3｜盲法验证机械阈值、发育窗口与迟滞",
        "claim": "主张：EFE-like 形成取决于力学波形、暴露时序和正反馈阈值；恢复平均血流并不总能逆转已锁定的状态。",
        "panels": [
            ("A", "保留扰动矩阵与实验随机化", "展示未参与模型拟合的扰动类型、强度、起止时间和恢复条件。", "至少设置三条相互正交的力学轴；按鱼随机分组、跨繁殖批次重复、成像和终点分析盲法。", "能够分别检验剪切、应变/压力、发育时序，而非把所有处理归为“低血流”。"),
            ("B", "匹配平均流量、改变脉动或相位", "两组具有相近平均流量/心输出，却有不同峰值、振荡指数或壁运动-流动相位差。", "通过不同操控组合实现物理量匹配；用 4D 成像确认实际暴露，而非仅依赖处理标签。", "两组产生不同 EFE-like 结局，直接支持“波形/相位比均值更关键”的反直觉预测。"),
            ("C", "模型生成个体化空间风险图", "在不知道组织学终点的情况下，为每条鱼输出心内膜表面的风险热图。", "冻结模型后输入该鱼的几何与力学历史；预先规定风险区面积和峰值位置指标。", "预测高风险区与后续 ECM/细胞状态异常空间共定位。"),
            ("D", "预测-实测空间配准", "风险热图与活体 ECM 信号、终点组织学及分子原位结果叠加。", "采用解剖坐标和表面距离度量，报告空间 AUC、峰值偏差及阴性区域特异性。", "模型不仅预测“会不会得病”，还能预测“在哪里首先发生”。"),
            ("E", "发育敏感窗口", "相同强度和持续时间的扰动在早、中、晚发育窗口产生不同结局。", "严格匹配处理剂量和恢复时间；同时监测整体发育阶段，区分时龄与心脏成熟度。", "只有模型预测的窗口跨越状态阈值，说明发育成熟改变了系统可塑性。"),
            ("F", "剂量阈值与非线性响应", "处理强度或机械剂量与 EFE-like 面积、厚度和功能之间的曲线。", "采用连续剂量而非二分类；比较线性、阈值和双稳态模型，给出不确定区间。", "终点出现陡峭转折而非均匀线性增加，与相图预测一致。"),
            ("G", "恢复实验揭示迟滞", "从相图不同位置撤除扰动，比较恢复后的轨迹；展示进入与退出阈值不重合。", "在阈值前、临界期和阈值后恢复物理条件；同一条鱼纵向追踪 ECM 和功能。", "阈值前可恢复、阈值后持续，支持 ECM-力学正反馈形成机械记忆。"),
            ("H", "盲预测成绩单与失败项", "用统一图表汇总空间、时间窗、剂量和恢复预测，不隐藏阴性预测。", "报告预测误差、校准曲线、批次外表现和预设成功标准；将模型修订放入扩展数据。", "至少两项关键盲预测成功，且成功不能由平均流量或一般发育延迟模型解释。"),
        ],
        "gate": "空间位置、发育窗口和恢复迟滞三类预测中至少两类在独立批次成立；否则停止扩展组学，先修订机制模型。",
    },
    {
        "title": "Figure 4｜解析细胞来源并把机械历史映射到细胞状态",
        "claim": "主张：EFE-like 并非单一来源成纤维细胞的静态堆积，而是来源和时间依赖的细胞状态转换；局部机械历史决定哪些谱系进入锁定状态。",
        "panels": [
            ("A", "心内膜与心外膜的正交谱系策略", "示意心内膜和心外膜/既有间充质细胞的独立诱导标记、标记时间和终点判定。", "使用经验证的心内膜谱系驱动线与 tcf21 等心外膜谱系系统或等效方案；设置漏标、未诱导和组织特异性对照。", "能够在同一疾病框架中定量比较不同来源，避免仅凭 marker 共表达推断来源。"),
            ("B", "活体追踪最早发生转换的细胞", "连续显示标记细胞形态、位置、迁移和 ECM 报告信号的变化。", "光片成像结合稀疏标记/光转换；记录从铺路石样心内膜到间充质样状态或 ECM 分泌状态的时间。", "确定细胞状态转换发生在 ECM 增厚之前还是之后，为因果方向提供时间证据。"),
            ("C", "不同来源对病变的时序贡献", "堆叠面积图或流图显示心内膜、心外膜和未定来源细胞在起始、扩展和维持阶段的比例。", "三维组织体积分数和细胞计数；以鱼为统计单位，细胞作为嵌套观测。", "检验“早期心内膜启动、后期心外膜/间充质扩增”等双来源模型，而不是强迫二选一。"),
            ("D", "机械状态分层的多组学采样设计", "四类样本：阈值前、转换中、锁定后、救援后；每类同时保留机械暴露和谱系信息。", "优先选择 scRNA+scATAC 或可行的多组学组合；控制发育期、扰动批次和细胞解离偏倚。", "采样围绕状态转换问题，不再制作缺少物理坐标的描述性全时程 atlas。"),
            ("E", "谱系注释的细胞状态图", "UMAP/状态图显示正常心内膜、激活内皮、过渡态、修复性成纤维细胞和锁定 fibroelastic 状态。", "结合谱系条码、经典标记、ECM 模块和细胞周期/应激评分；避免仅依赖拟时序判断来源。", "识别一个与持续病变和纤维弹性 ECM 特异相关的细胞状态。"),
            ("F", "状态转换方向与分支", "谱系已知条件下的转录/染色质轨迹，比较回归正常与进入锁定态两条分支。", "使用 RNA velocity/轨迹方法仅作辅助，主结论由真实谱系和纵向成像约束。", "锁定态不是一般激活程度更高，而是具有独立调控程序和回不到正常态的分支。"),
            ("G", "细胞状态与局部机械剂量的空间耦合", "将原位验证的细胞状态投影到 Figure 2-3 的机械风险图，展示逐区域关系。", "用多重原位/HCR、空间转录或区域光标记保持空间信息；采用混合效应模型控制个体和批次。", "机械剂量可解释同一心脏内细胞状态异质性，并优于解剖位置或平均流量。"),
            ("H", "调控网络筛选 Gate X", "展示从差异可及性、转录因子活性、配体-受体和模型敏感性共同筛出的候选门控节点。", "优先考察与力学感知、TGFβ/BMP 平衡、黏附/ECM 反馈有关的节点；以跨数据层一致性排序。", "得到一个适合细胞特异、时间可控干预的主节点 Gate X，并预先排除只反映终末纤维化的旁观者。"),
            ("I", "来源争议的统一时序模型", "总结图显示不同来源细胞在不同机械区间和时间阶段的作用。", "用谱系、活体成像、空间组学和理论风险图共同约束，不依赖单一技术。", "提出能同时容纳 EndMT 证据与心外膜来源证据的时序机制，并生成 Figure 5 的因果检验。"),
        ],
        "gate": "必须锁定一个与持续状态相关、能被空间验证且具有可操作性的 Gate X；若只有非特异炎症/应激信号，不进入大规模干预。",
    },
    {
        "title": "Figure 5｜证明机械-分子门控的必要性、充分性与可逆性",
        "claim": "主张：Gate X 将机械暴露转译为持续的 fibroelastic 状态；细胞特异且时间正确的干预可以移动相变阈值并恢复心功能。",
        "panels": [
            ("A", "模型选定的机制轴", "将 Figure 4 得到的 Gate X 放入心内膜状态-ECM-流动反馈环，标出预测的上游和下游。", "整合力学传感、TGFβ/BMP 或黏附/ECM 网络；主图只保留一条具有最强因果证据的轴。", "给出明确的遗传互作预测：阻断何处可预防，何处可逆转，何处只能改善终点。"),
            ("B", "Gate X 的时空激活先于病变", "Gate X 报告、核定位或靶基因活动与机械风险图、最早 ECM 沉积的叠加。", "活体 reporter 或定量免疫/原位；高时间分辨率采样，比较风险高区与同心脏风险低区。", "Gate X 激活发生在细胞状态转换和 ECM 增厚之前，并随机械剂量变化。"),
            ("C", "谱系特异必要性", "分别在心内膜和心外膜/间充质谱系抑制 Gate X，比较病变形成。", "诱导性、组织特异的 CRISPR/转基因抑制或等效手段；验证编辑效率和非目标组织功能。", "只有关键谱系中的抑制显著降低锁定细胞状态和 EFE-like ECM，确定主要作用位置。"),
            ("D", "在正常力学环境中的充分性", "正常负荷下激活 Gate X，观察是否能部分重现细胞状态和 ECM 表型。", "使用低水平、时间限制的细胞特异激活，避免超生理表达和广泛毒性。", "Gate X 激活足以推动核心状态，但若仍需机械共同作用，则形成更精确的 AND-gate 机制。"),
            ("E", "遗传上位性与通路顺序", "上游力学感知节点、Gate X 和 ECM 执行节点的双重扰动矩阵。", "比较单独与联合干预，对细胞状态、ECM 和功能进行一致读出。", "确定 Gate X 位于机械输入与 ECM 锁定之间，而非疾病形成后的伴随反应。"),
            ("F", "阈值前干预的预防性救援", "在模型预测的阈值前短暂抑制 Gate X，随后不再用药/诱导，观察长期结局。", "严格匹配干预窗口；同时记录机械暴露，证明救援不是因为改变了心率或整体血流。", "短暂早期干预即可阻止进入锁定态，验证发育窗口。"),
            ("G", "阈值后干预与反馈拆解", "在已形成 fibroelastic layer 后分别恢复血流、抑制 Gate X、干预 ECM 反馈，比较逆转能力。", "设置机械恢复、分子干预、ECM 干预及组合组；纵向随访而非单终点。", "组合或反馈端干预能够使系统跨回可逆区，解释为何单纯恢复血流可能失败。"),
            ("H", "结构、细胞状态与心功能同步救援", "展示救援后 ECM 厚度、锁定细胞状态比例、顺应性和舒张功能的多层一致恢复。", "采用与 Figure 1 相同的预注册终点和盲法分析；报告完全、部分和无响应个体。", "干预不仅降低 marker，还恢复组织结构和心功能。"),
            ("I", "干预使相图边界发生可预测移动", "把干预组重新投影到 Figure 2 的相图，显示阈值位置或迟滞环缩小。", "在不重新拟合全部参数的条件下只更新 Gate X 相关参数，并预测新的盲测条件。", "分子因果结果与理论参数同向变化，完成模型-实验闭环。"),
        ],
        "gate": "同一干预必须同时改变细胞状态、fibroelastic ECM 和心功能，并能由模型解释为阈值移动；只改善一个终点不足以进入转化主图。",
    },
    {
        "title": "Figure 6｜在人 EFE 中验证同一状态并重建可干预机制",
        "claim": "主张：斑马鱼发现的机械记忆与 Gate X 不是物种特异现象，而是人 EFE 中可识别、可在受控人源体系中重现并干预的疾病机制。",
        "panels": [
            ("A", "人 EFE 样本与对照框架", "展示 EFE/HLHS 病例、非 EFE 先心病对照和必要的年龄/解剖匹配信息。", "第一年启动儿科心脏外科、病理和伦理合作；预先定义纳入标准、组织区域和临床影像变量。", "样本体系能够区分 EFE 特异改变、一般心衰纤维化和手术/年龄效应。"),
            ("B", "人组织的三维 fibroelastic 病理", "连续切片或体成像重建心内膜下胶原/弹性层，并量化厚度、取向和空间异质性。", "组织学、免疫/原位和 ECM 成分分析；与术前功能/流动指标在可获得范围内关联。", "人病变在结构和 ECM 组成上与斑马鱼 EFE-like 状态具有可比特征。"),
            ("C", "人 EFE 的细胞状态与 Gate X 活性", "空间转录、snRNA/scRNA 或多重原位显示锁定 fibroelastic 状态和 Gate X 网络。", "考虑冷冻/固定样本差异，优先使用能保存病变空间位置的方法；设置心内膜、心外膜和间充质来源标记组合。", "在人病变核心区发现与鱼锁定态同源的调控模块，而非只看到共有 fibrosis genes。"),
            ("D", "跨物种状态映射", "展示鱼和人细胞状态的同源基因模块、调控子和 ECM 程序映射。", "使用一对一/多对多同源基因处理、伪 bulk 和个体层验证；留出独立病例作为验证集。", "Gate X 和关键下游模块在跨物种映射中稳定，效应不由单个 marker 驱动。"),
            ("E", "人源心内膜力学培养系统", "示意人 iPSC 来源或原代心内膜样细胞在可控流动、周期应变和基质刚度下培养。", "输入 Figure 2 定义的物理变量，建立正常、阈值附近和锁定区三类条件；记录表型与分泌 ECM。", "人源细胞在模型预测的组合负荷下进入 EFE-like 状态，而非任意高应力均产生同样结果。"),
            ("F", "重现阈值、时间窗与机械记忆", "在人源体系中展示剂量阈值、暴露时长和撤除刺激后的持续性。", "采用独立批次和多条 iPSC 系；把人源实验参数按无量纲或可比物理量映射到模型。", "系统重现非线性状态转换和迟滞，强化普适机制而非物种类比。"),
            ("G", "Gate X 的人源救援", "遗传或药理干预 Gate X，观察细胞状态、ECM 和力学表型恢复。", "使用机制特异干预并监测细胞毒性、屏障功能和一般内皮反应；与既有候选治疗作参照而非替代主机制。", "同一节点在鱼和人源体系均产生方向一致的救援。"),
            ("H", "统一模型与临床可检验预测", "最终示意异常力学历史如何启动心内膜状态转换、ECM 正反馈、迟滞和功能受限，并标出干预窗口。", "用全部数据冻结模型，输出可供后续胎儿超声/组织研究检验的指标，如波形、时间窗或空间风险。", "论文结论从“一个斑马鱼通路”上升为发育期纤维化不可逆化的普适机制。"),
        ],
        "gate": "至少在人 EFE 组织中验证锁定状态/Gate X，并在人源或哺乳动物体系完成一次因果重现或救援，才支撑 Nature 级跨物种结论。",
    },
]


EXTENDED_DATA = [
    "Extended Data 1：三类扰动的存活率、整体发育、心率、炎症和非特异毒性控制。",
    "Extended Data 2：4D 分割、表面重建、网格质量和跨操作者重复性。",
    "Extended Data 3：EFE-like ECM 的多种正交染色/蛋白/转录验证及阴性区域。",
    "Extended Data 4：流体与组织力学模型的收敛、敏感性、参数可辨识性和替代模型比较。",
    "Extended Data 5：谱系工具的诱导效率、漏标、组织特异性和追踪稳定性。",
    "Extended Data 6：单细胞/多组学 QC、批次效应、细胞数、个体重复和完整候选通路。",
    "Extended Data 7：Gate X 的剂量、时间、组织特异性、脱靶和一般内皮功能控制。",
    "Extended Data 8：人样本临床信息、病理分层、样本质量、个体层统计和独立验证病例。",
]


REFERENCES = [
    "Endocardial Fibroelastosis is Caused by Aberrant Endothelial to Mesenchymal Transition. https://pmc.ncbi.nlm.nih.gov/articles/PMC4344885/",
    "Fibroblasts in an endocardial fibroelastosis disease model mainly originate from mesenchymal derivatives of epicardium. https://pubmed.ncbi.nlm.nih.gov/28809397/",
    "Abnormal Flow Conditions Promote Endocardial Fibroelastosis Via Endothelial-to-Mesenchymal Transition, Which Is Responsive to Losartan Treatment. https://pmc.ncbi.nlm.nih.gov/articles/PMC8733675/",
    "Mechanical strain triggers endothelial-to-mesenchymal transition of the endocardium in the immature heart. https://pmc.ncbi.nlm.nih.gov/articles/PMC9133271/",
    "Origin and function of activated fibroblast states during zebrafish heart regeneration. https://www.nature.com/articles/s41588-022-01129-5",
    "Transient fibrosis resolves via fibroblast inactivation in the regenerating zebrafish heart. https://pmc.ncbi.nlm.nih.gov/articles/PMC5910827/",
    "An organ-wide spatiotemporal transcriptomic and cellular atlas of the regenerating zebrafish heart. https://www.nature.com/articles/s41467-025-59070-0",
]


def build_document():
    doc = Document()
    setup_styles(doc)
    configure_sections(doc)

    section = doc.sections[0]
    section.different_first_page_header_footer = False

    # Editorial cover
    for _ in range(5):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run("研究主图设计 · 3–4 年路线")
    set_run_font(r, size=11, color=GOLD, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("心内膜弹性纤维增生症（EFE）")
    set_run_font(r, size=28, color=INK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(30)
    r = p.add_run("面向 Nature 的六张主图与逐 Panel 实验设计")
    set_run_font(r, size=16, color=DARK_BLUE, bold=False)

    add_callout(
        doc,
        "中心命题：发育期异常血流形成可累积的机械记忆，通过心内膜-ECM 正反馈，将可逆修复性纤维化锁定为不可逆的 EFE。",
        fill=CALL_OUT,
        border="2E74B5",
    )

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(42)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("优势组合")
    set_run_font(r, size=10, color=MUTED, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(40)
    r = p.add_run("斑马鱼活体成像 × 三维/四维重建 × 理论力学 × 生信与跨物种分析")
    set_run_font(r, size=11.5, color=INK, bold=True)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("内部研究设计稿｜2026年8月")
    set_run_font(r, size=9.5, color=MUTED)

    doc.add_page_break()

    doc.add_heading("论文叙事骨架", level=1)
    add_callout(doc, "现象发现（4D成像）→ 个体化重建（力学历史）→ 盲预测（相图与阈值）→ 细胞来源（谱系）→ 分子门控（多组学与因果）→ 人类疾病验证。")

    doc.add_heading("术语和边界", level=2)
    for text in [
        "在斑马鱼完成跨物种验证前，使用“EFE-like fibroelastic remodeling”，不直接宣称建立了人 EFE。",
        "疾病表型必须同时包含心内膜下定位、纤维弹性 ECM、持续性和功能后果；仅有 col1a/postn 上升不构成 EFE。",
        "成年斑马鱼损伤/再生模型主要作为“可逆修复性纤维化”参照，主疾病模型应放在发育期心脏。",
        "细胞来源允许为时间依赖的双来源或多来源，不预设 EndMT 与心外膜来源必须二选一。",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("全项目统计和证据原则", level=2)
    for text in [
        "鱼是生物学重复单位；图像块、切片、细胞和材料点属于嵌套观测，不得充当独立 n。",
        "模型校准组与盲测组在看到病理终点前划分；主要预测、误差阈值和失败标准预先登记。",
        "每个关键结论至少由两种正交技术支持，例如活体信号 + 终点组织学、谱系 + 空间原位、结构救援 + 功能救援。",
        "主图只保留支持中心命题的结果；完整 QC、阴性结果、替代模型和通路筛选进入 Extended Data。",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("六张主图的一句话功能", level=2)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    set_table_fixed_geometry(table, [1650, 7710])
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    headers = ["主图", "该图必须回答的问题"]
    for i, text in enumerate(headers):
        hdr.cells[i].text = text
        set_cell_shading(hdr.cells[i], LIGHT_BLUE)
        for run in hdr.cells[i].paragraphs[0].runs:
            set_run_font(run, size=10.5, color=INK, bold=True)
    rows = [
        ("Figure 1", "我们是否真正建立了具有疾病定义的 EFE-like 表型？"),
        ("Figure 2", "能否从个体 4D 图像计算局部机械历史并提出可证伪预测？"),
        ("Figure 3", "模型是否在未见数据上预测空间、时间窗、阈值和迟滞？"),
        ("Figure 4", "哪些谱系在什么时间进入锁定细胞状态，机械历史如何解释异质性？"),
        ("Figure 5", "Gate X 是否具有必要性、充分性，并可被时间特异地救援？"),
        ("Figure 6", "同一状态和机制是否存在于人 EFE，并可在人源体系中重现？"),
    ]
    for label, value in rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value
        for idx, cell in enumerate(cells):
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    set_run_font(run, size=10.2, color=INK if idx == 0 else None, bold=(idx == 0))

    for fig in FIGURES:
        doc.add_page_break()
        doc.add_heading(fig["title"], level=1)
        add_callout(doc, fig["claim"])
        for panel in fig["panels"]:
            add_panel(doc, *panel)
        add_gate(doc, fig["gate"])

    doc.add_page_break()
    doc.add_heading("Extended Data 分配建议", level=1)
    add_callout(doc, "原则：主图回答机制问题，Extended Data 证明结论可靠、边界清楚、替代解释被排除。")
    for item in EXTENDED_DATA:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("3–4 年并行排程建议", level=1)
    schedule = doc.add_table(rows=1, cols=4)
    schedule.style = "Table Grid"
    set_table_fixed_geometry(schedule, [1080, 2700, 2700, 2880])
    headers = ["阶段", "成像/实验", "理论/生信", "阶段闸门"]
    for i, text in enumerate(headers):
        cell = schedule.rows[0].cells[i]
        cell.text = text
        set_cell_shading(cell, LIGHT_BLUE)
        for run in cell.paragraphs[0].runs:
            set_run_font(run, size=9.5, color=INK, bold=True)
    set_repeat_table_header(schedule.rows[0])
    schedule_rows = [
        ("第1年", "建立 EFE-like 模型；完成纵向4D成像、ECM与功能判定。", "最小状态模型；材料坐标和重建流水线。", "Figure 1 四项疾病标准成立。"),
        ("第2年", "完成正交扰动、恢复实验和双谱系先导。", "相图冻结；完成盲预测与空间风险图。", "Figure 3 至少两类预测成立。"),
        ("第3年", "完成谱系、多组学、Gate X 因果和早/晚救援。", "力学-状态调控网络；模型参数与干预闭环。", "结构、状态、功能三层同步救援。"),
        ("第4年", "完成人组织/人源体系验证和独立重复。", "跨物种映射；最终模型冻结与数据发布。", "Figure 6 支撑普适机制。"),
    ]
    for row_data in schedule_rows:
        cells = schedule.add_row().cells
        for i, text in enumerate(row_data):
            cells[i].text = text
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.1
                for run in p.runs:
                    set_run_font(run, size=8.9, color=INK if i == 0 else None, bold=(i == 0))

    doc.add_heading("最早需要落实的外部能力", level=2)
    for text in [
        "儿科心脏外科/病理合作：EFE/HLHS 组织、非 EFE 先心病对照、伦理和临床信息。",
        "稳定的谱系遗传工具与细胞特异、时间可控的因果干预能力。",
        "人 iPSC 来源心内膜或原代心内膜合作，以及可控制流动、应变和基质刚度的体外体系。",
        "必要时引入局部组织力学测量合作，但其方法学必须服务于 Figure 2 的可测模型参数。",
    ]:
        doc.add_paragraph(text, style="List Bullet")

    doc.add_page_break()
    doc.add_heading("工作参考文献", level=1)
    p = doc.add_paragraph(style="Panel Body")
    p.add_run("以下文献用于确定 EFE 细胞来源争议、异常流动/机械应变机制，以及斑马鱼纤维化-再生的当前竞争基线。")
    for ref in REFERENCES:
        p = doc.add_paragraph(ref, style="List Number")
        for run in p.runs:
            set_run_font(run, size=9.5)

    # Footer label and core properties.
    props = doc.core_properties
    props.title = "EFE Nature 主图小图设计方案"
    props.subject = "心内膜弹性纤维增生症；斑马鱼；4D成像；理论模型；生信；Nature主图设计"
    props.author = "EFE research planning team"
    props.keywords = "EFE, zebrafish, mechanobiology, fibrosis, imaging, computational model"

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    path = build_document()
    print(path)
